import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from database import get_all_attendance_records

def generate_attendance_excel(start_date=None, end_date=None, department=None, emp_id=None, output_path=None):
    records = get_all_attendance_records(start_date=start_date, end_date=end_date, department=department, emp_id=emp_id)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Report"

    # Header Styling
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Title Block
    ws.merge_cells('A1:H1')
    title_cell = ws['A1']
    title_cell.value = "SMART ATTENDANCE MANAGEMENT SYSTEM - ATTENDANCE REPORT"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="1E293B")
    title_cell.alignment = align_center
    ws.row_dimensions[1].height = 30

    # Subtitle with metadata
    ws.merge_cells('A2:H2')
    sub_cell = ws['A2']
    gen_time = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    sub_cell.value = f"Generated on: {gen_time} | Filter: Dept={department or 'All'}, Range={start_date or 'All'} to {end_date or 'All'}"
    sub_cell.font = Font(name="Calibri", size=10, italic=True, color="64748B")
    sub_cell.alignment = align_center
    ws.row_dimensions[2].height = 20

    headers = [
        "Record ID", "Employee ID", "Employee Name", "Department",
        "Date", "Punch IN Time", "Punch OUT Time", "Total Hours", "Status"
    ]

    # Write Column Headers
    header_row = 4
    ws.row_dimensions[header_row].height = 25
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = thin_border

    # Write Data Rows
    row_num = header_row + 1
    for r in records:
        ws.row_dimensions[row_num].height = 22
        
        status_text = r.get('status', 'IN_PROGRESS')
        if status_text == 'COMPLETED':
            status_display = 'Completed (Out)'
        elif status_text == 'IN_PROGRESS':
            status_display = 'In Office (Active)'
        else:
            status_display = status_text

        row_data = [
            r['id'],
            r['emp_id'],
            r.get('name', ''),
            r.get('department', ''),
            r['date'],
            r['in_time'] or '-',
            r['out_time'] or '-',
            f"{r['total_hours']} hrs" if r['total_hours'] else '-',
            status_display
        ]

        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.alignment = align_center if col_idx in [1, 2, 5, 6, 7, 8, 9] else align_left
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)

        row_num += 1

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    if not output_path:
        filename = f"Attendance_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        output_path = os.path.join(os.path.dirname(__file__), 'captures', filename)

    wb.save(output_path)
    return output_path

def generate_attendance_csv(start_date=None, end_date=None, department=None, emp_id=None, output_path=None):
    import csv
    records = get_all_attendance_records(start_date=start_date, end_date=end_date, department=department, emp_id=emp_id)
    if not output_path:
        filename = f"Attendance_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        output_path = os.path.join(os.path.dirname(__file__), 'captures', filename)

    with open(output_path, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(["Record ID", "Employee ID", "Employee Name", "Department", "Date", "Punch IN Time", "Punch OUT Time", "Total Hours", "Status"])
        for r in records:
            status_text = r.get('status', 'IN_PROGRESS')
            status_display = 'Completed (Out)' if status_text == 'COMPLETED' else ('In Office (Active)' if status_text == 'IN_PROGRESS' else status_text)
            writer.writerow([
                r['id'],
                r['emp_id'],
                r.get('name', ''),
                r.get('department', ''),
                r['date'],
                r['in_time'] or '-',
                r['out_time'] or '-',
                f"{r['total_hours']} hrs" if r['total_hours'] else '-',
                status_display
            ])
    return output_path
