import cv2
import numpy as np
import pickle
import os
from datetime import datetime
from mtcnn.mtcnn import MTCNN
from keras_facenet import FaceNet
import openpyxl
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage
import io

# 1. Configuration
MODEL_PATH = 'face_classification_model.pkl'
ENCODER_PATH = 'label_encoder.pkl'
DATASET_PATH = 'Dataset'
ATTENDANCE_FILE = f'Attendance_{datetime.now().strftime("%Y-%m-%d")}.xlsx'

print("Loading MTCNN and FaceNet...")
detector = MTCNN()
embedder = FaceNet()

if not os.path.exists(MODEL_PATH) or not os.path.exists(ENCODER_PATH):
    print("Error: Model files not found! Please run 1_train_model.py first.")
    exit()

print("Loading trained classifier...")
with open(MODEL_PATH, 'rb') as f:
    model = pickle.load(f)
with open(ENCODER_PATH, 'rb') as f:
    encoder = pickle.load(f)

marked_students = set()

def get_reference_photo(name):
    """Finds the first photo of the student and returns its path."""
    person_dir = os.path.join(DATASET_PATH, name)
    if os.path.exists(person_dir):
        images = os.listdir(person_dir)
        if len(images) > 0:
            return os.path.join(person_dir, images[0])
    return None

# Initialize the Excel file with new headers
if not os.path.exists(ATTENDANCE_FILE):
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"
    student_names = list(encoder.classes_)
    
    for i, name in enumerate(student_names):
        col = i + 1
        ws.cell(row=1, column=col, value=name)
        
        # Add reference photo to the second row
        ref_path = get_reference_photo(name)
        if ref_path:
            try:
                img = XLImage(ref_path)
                img.width = 100
                img.height = 100
                ws.add_image(img, f"{get_column_letter(col)}2")
            except Exception as e:
                print(f"Warning: Could not load reference image for {name}")

    ws.row_dimensions[2].height = 80
    wb.save(ATTENDANCE_FILE)

def mark_attendance(name, face_img):
    """Marks attendance and saves the captured face image in the Excel file."""
    if name not in marked_students and name != "Unknown":
        now = datetime.now()
        time_string = now.strftime('%H:%M:%S')
        
        try:
            wb = load_workbook(ATTENDANCE_FILE)
            ws = wb.active
            student_names = list(encoder.classes_)
            
            if name in student_names:
                col_idx = student_names.index(name) + 1
                row_idx = ws.max_row + 1
                
                # Write timestamp
                ws.cell(row=row_idx, column=col_idx, value=time_string)
                
                # Process and add captured face image
                pil_img = PILImage.fromarray(face_img)
                img_io = io.BytesIO()
                pil_img.save(img_io, format='JPEG')
                img_io.seek(0)
                
                img = XLImage(img_io)
                img.width = 100
                img.height = 100
                
                ws.add_image(img, f"{get_column_letter(col_idx)}{row_idx}")
                ws.row_dimensions[row_idx].height = 80
                
                wb.save(ATTENDANCE_FILE)
                marked_students.add(name)
                print(f"✅ Attendance marked for: {name} at {time_string}")
        except Exception as e:
            print(f"Error marking attendance: {e}")

def extract_face(frame, box, required_size=(160, 160)):
    x, y, width, height = box
    x1, y1 = abs(x), abs(y)
    x2, y2 = x1 + width, y1 + height
    face_pixels = frame[y1:y2, x1:x2]
    return cv2.resize(face_pixels, required_size)

# 2. Start Live Camera
print("Starting live camera feed. Press 'q' to quit.")
cap = cv2.VideoCapture(0) 

while True:
    ret, frame = cap.read()
    if not ret:
        print("Camera not found.")
        break

    rgb_frame = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
    results = detector.detect_faces(rgb_frame)

    for result in results:
        if result['confidence'] < 0.90:
            continue
            
        box = result['box']
        face = extract_face(rgb_frame, box)
        
        # Get embeddings and predict
        face_array = np.expand_dims(face, axis=0)
        embeddings = embedder.embeddings(face_array)
        
        prediction = model.predict(embeddings)
        probability = model.predict_proba(embeddings)
        confidence = np.max(probability)
        
        # Using 0.50 as the threshold as discussed earlier!
        if confidence > 0.50: 
            class_index = prediction[0]
            name = encoder.inverse_transform([class_index])[0]
            # Pass the extracted face to the attendance function
            mark_attendance(name, face)
        else:
            name = "Unknown"

        # Draw UI
        x, y, w, h = box
        color = (0, 255, 0) if name != "Unknown" else (0, 0, 255)
        cv2.rectangle(frame, (x, y), (x+w, y+h), color, 2)
        display_text = f"{name} ({confidence*100:.1f}%)" if name != "Unknown" else "Unknown"
        cv2.putText(frame, display_text, (x, y - 10), cv2.FONT_HERSHEY_SIMPLEX, 0.8, color, 2)

    cv2.imshow('Smart Classroom - Live Attendance', frame)

    if cv2.waitKey(1) & 0xFF == ord('q'):
        break

cap.release()
cv2.destroyAllWindows()